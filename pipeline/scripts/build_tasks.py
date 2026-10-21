# 직업별 '세부 작업'(O*NET DWA) 목록을 만든다 → tasks_raw.json
# 소분류(3자리)·중분류(2자리)마다: 연결된 미국 직업(SOC)의 업무 문장 → 세부 작업(DWA)으로 묶고,
# 업무 중요도(IM, 없으면 같은 무게)·핵심/보조(Core 1, Supplemental 0.5)·미국 고용으로 가중해 상위 작업을 고른다.
# 소식은 이 세부 작업 단위로 분류하고(feed.json의 dwa), 직업 카드는 이 목록으로 '이 발표가 닿는 내 업무'를 보여준다.
import csv, json, collections
exec(open('soc_map_small.py').read()); exec(open('soc_map.py').read())
O = 'dl_onet/db_30_0_text/'
rd = lambda f: csv.DictReader(open(O + f, encoding='utf-8'), delimiter='\t')
EMP = {}
for r in csv.DictReader(open('gpts/national_May2021_dl.csv', encoding='utf-8-sig')):
    if r['O_GROUP'] == 'detailed':
        try: EMP[r['OCC_CODE']] = float(r['TOT_EMP'].replace(',', ''))
        except: pass
def emp(soc): return EMP.get(soc) or next((x for k, x in EMP.items() if k[:6] == soc[:6]), 1000.0)
TT = {}; TS = {}
for r in rd('Task Statements.txt'):
    TT[(r['O*NET-SOC Code'], r['Task ID'])] = r['Task Type'] or 'Core'; TS[r['Task ID']] = r['Task']
IM = {(r['O*NET-SOC Code'], r['Task ID']): float(r['Data Value']) for r in rd('Task Ratings.txt') if r.get('Scale ID', 'IM') == 'IM'}
T2D = collections.defaultdict(list)
for r in rd('Tasks to DWAs.txt'): T2D[(r['O*NET-SOC Code'], r['Task ID'])].append(r['DWA ID'])
DT = {r['DWA ID']: r['DWA Title'] for r in rd('DWA Reference.txt')}
# 직업(O*NET 코드)별 세부 작업 몫
occ = collections.defaultdict(dict)
for (o, t), typ in TT.items():
    ds = T2D.get((o, t))
    if not ds: continue
    w = IM.get((o, t), 3.0) * (1 if typ == 'Core' else .5)
    for d in ds: occ[o][d] = occ[o].get(d, 0) + w / len(ds)
groups = {'sm': collections.defaultdict(lambda: collections.Counter()), 'mid': collections.defaultdict(lambda: collections.Counter()), 'ks': collections.defaultdict(lambda: collections.Counter())}
for o, dw in occ.items():
    soc = o[:7]; tot = sum(dw.values()); e = emp(soc)
    for key, f in (('sm', small_of), ('mid', ksco_of)):
        c = f(soc)
        if not c: continue
        for d, v in dw.items(): groups[key][c][d] += e * v / tot
# 세분류(4자리): 국가데이터처 KSCO↔ISCO 공식 연계표 → ISCO ↔ O*NET-SOC(O*NET의 ESCO 연계표, ESCO 코드 앞 4자리 = ISCO) 
import openpyxl
xw = collections.defaultdict(collections.Counter)
for i, r in enumerate(openpyxl.load_workbook('onet_xw/ESCO_to_ONET-SOC.xlsx', read_only=True).worksheets[0].iter_rows(values_only=True)):
    if i >= 4 and r[0] and r[2]: xw[str(r[0])[:4]][r[2]] += 1
ks2i = collections.defaultdict(list)
for i, r in enumerate(openpyxl.load_workbook('ksco_isco/KSCO8_ISCO08_20251103.xlsx', read_only=True)['4-1. (연계표) 세분류 연계표(KSCO-ISCO)'].iter_rows(values_only=True)):
    if i and r[0] and str(r[0])[:1].isdigit(): ks2i[str(r[0])].append(str(r[4]).lstrip('*').zfill(4))
KSO = {}
for k, isos in ks2i.items():
    isos = list(dict.fromkeys(isos)); w = collections.Counter()
    for x in isos:
        tot = sum(xw[x].values())
        for o, n in xw[x].items(): w[o] += n / tot / len(isos)
    KSO[k] = w
    for o, s in w.items():
        dw = occ.get(o)
        if not dw: continue
        t = sum(dw.values())
        for d, v in dw.items(): groups['ks'][k][d] += s * v / t
OT = {r['O*NET-SOC Code']: r['Title'] for r in rd('Occupation Data.txt')}
def top(cnt):
    tot = sum(cnt.values()); acc = 0; pick = []
    for d, v in cnt.most_common():
        if len(pick) >= 12 or (len(pick) >= 6 and acc >= .6): break
        pick.append([d, v / tot]); acc += v / tot
    s = sum(x for _, x in pick); return [[d, round(x / s, 4)] for d, x in pick]
out = {'dwa': {}, 'sm': {}, 'mid': {}, 'ks': {}, 'occ': {}}
# 세분류마다 '세부 직무' 후보(O*NET 직업) — 몫 4% 이상, 최대 6개. 사용자가 가장 가까운 것을 고르면 그 직무의 세부 작업을 보여준다.
# 연계표만으로 고른 후보는 다른 분야가 섞인다(예: 로봇공학 기술자 → 태양광 엔지니어).
# 세분류마다 사람이 검토한 목록(ks_onet_curated.json)이 있으면 그것을 쓴다. 첫 번째가 기본 직무.
import os
CUR = json.load(open('ks_onet_curated.json')) if os.path.exists('ks_onet_curated.json') else {}
for k in (set(CUR) if CUR else set(KSO)):  # 검토 목록이 있으면 그 세분류만(직업명 사전에 없는 코드는 화면에서 쓰이지 않음)
    if k in CUR:
        c = [[o, s] for o, s in CUR[k]['o'] if o in occ]
        t = sum(s for _, s in c) or 1; c = [[o, round(s / t, 3)] for o, s in c]
    else:
        w = KSO[k]; tot = sum(v for o, v in w.items() if o in occ) or 1
        c = [[o, round(v / tot, 3)] for o, v in w.most_common() if o in occ and v / tot >= .04][:6]
    if c: out['ks'][k] = c
    for o, _ in c:
        if o not in out['occ']:
            out['occ'][o] = {'t': OT.get(o, o), 'd': top(collections.Counter(occ[o]))}
            for d, _ in out['occ'][o]['d']: out['dwa'][d] = DT[d]
for key in ('sm', 'mid'):
    for c, cnt in groups[key].items():
        tot = sum(cnt.values()); acc = 0; pick = []
        for d, v in cnt.most_common():
            if len(pick) >= 12 or (len(pick) >= 6 and acc >= .6): break
            pick.append([d, round(v / tot, 4)]); acc += v / tot
        s = sum(x for _, x in pick); out[key][c] = [[d, round(x / s, 4)] for d, x in pick]
        for d, _ in pick: out['dwa'][d] = DT[d]
json.dump(out, open('tasks_raw.json', 'w'), ensure_ascii=False, indent=0)
print('ks', len(out['ks']), 'occ', len(out['occ']), 'sm', len(out['sm']), 'mid', len(out['mid']), 'dwa', len(out['dwa']))
# 사이트용 tasks.json: 한국어 이름(tasks_ko/)을 붙인다. 번역 파일은 처음 한 번 만들어 고정본으로 둔다.
import os
if os.path.exists('tasks_ko/dwa_ko.json') and os.path.exists('tasks_ko/occ_ko.json'):
    DK = json.load(open('tasks_ko/dwa_ko.json')); OK_ = json.load(open('tasks_ko/occ_ko.json'))
    miss = [d for d in out['dwa'] if d not in DK] + [o for o in out['occ'] if o not in OK_]
    if miss: raise SystemExit(f'번역 없는 항목 {len(miss)}개: {miss[:5]} — tasks_ko/에 더해야 함')
    site = {'v': 1, 'dwa': {d: DK[d] for d in out['dwa']}, 'occ': {o: [OK_[o], v['d']] for o, v in out['occ'].items()},
            'ks': out['ks'], 'sm': out['sm'], 'mid': out['mid']}
    json.dump(site, open('tasks.json', 'w'), ensure_ascii=False, separators=(',', ':'))
    print('tasks.json', os.path.getsize('tasks.json'))
